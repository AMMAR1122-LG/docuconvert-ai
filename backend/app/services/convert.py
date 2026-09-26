"""
Cross-format conversion. PDF -> Word uses pdf2docx directly (pure
Python, works in-process). Word/PPT/Excel -> PDF shells out to
LibreOffice headless, which is the standard production-grade approach
for faithful DOC/DOCX/PPTX/XLSX rendering — it must be installed in the
container (see docker/backend.Dockerfile).
"""
from __future__ import annotations

import io
import shutil
import subprocess
import tempfile
from pathlib import Path

import fitz  # PyMuPDF
from pdf2docx import Converter as Pdf2DocxConverter

from app.core.errors import AppError


def pdf_to_docx(data: bytes) -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in.pdf"
        dst = Path(tmp) / "out.docx"
        src.write_bytes(data)
        try:
            converter = Pdf2DocxConverter(str(src))
            converter.convert(str(dst))
            converter.close()
        except Exception as exc:
            raise AppError(
                "We couldn't convert this PDF to Word. It may be a scanned/image-only PDF — try OCR PDF first.",
                400,
                "conversion_failed",
            ) from exc
        return dst.read_bytes()


def _libreoffice_available() -> bool:
    return shutil.which("soffice") is not None or shutil.which("libreoffice") is not None


def _run_libreoffice(src: Path, out_dir: Path, target_filter: str) -> Path:
    binary = shutil.which("soffice") or shutil.which("libreoffice")
    if not binary:
        raise AppError(
            "Document conversion is temporarily unavailable. Please try again shortly.",
            503,
            "converter_unavailable",
        )
    cmd = [binary, "--headless", "--norestore", "--convert-to", target_filter, "--outdir", str(out_dir), str(src)]
    try:
        result = subprocess.run(cmd, capture_output=True, timeout=120)
    except subprocess.TimeoutExpired as exc:
        raise AppError("Conversion took too long. Try a smaller file.", 504, "conversion_timeout") from exc
    if result.returncode != 0:
        raise AppError("We couldn't convert this file. It may be corrupted or password-protected.", 400, "conversion_failed")
    # target_filter can carry options like "pdf:writer_pdf_Export" — the produced
    # file's extension is just the part before any colon. Match ONLY that
    # extension so we never pick up the original input file sitting in the
    # same directory (e.g. "in.docx" alongside the real output "in.pdf").
    out_ext = target_filter.split(":")[0]
    produced = list(out_dir.glob(f"{src.stem}.{out_ext}"))
    if not produced:
        raise AppError("Conversion produced no output.", 500, "conversion_failed")
    return produced[0]


def docx_to_pdf(data: bytes, original_suffix: str = ".docx") -> bytes:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        src = tmp_path / f"in{original_suffix}"
        src.write_bytes(data)
        out_path = _run_libreoffice(src, tmp_path, "pdf")
        result = out_path.read_bytes()
        if not result.startswith(b"%PDF-"):
            raise AppError("Conversion produced an invalid PDF.", 500, "conversion_failed")
        return result


def pdf_to_pptx_fallback(data: bytes) -> bytes:
    """
    Page-based fallback per spec section 13: each PDF page becomes a
    full-bleed image on its own slide. This always works, unlike
    text-reflow conversion which is unreliable for arbitrary PDFs.
    """
    from pptx import Presentation
    from pptx.util import Emu

    doc = fitz.open(stream=data, filetype="pdf")
    prs = Presentation()
    prs.slide_width = Emu(12192000)  # 16:9 widescreen
    prs.slide_height = Emu(6858000)
    blank_layout = prs.slide_layouts[6]

    for page in doc:
        pix = page.get_pixmap(dpi=150)
        img_bytes = pix.tobytes("png")
        slide = prs.slides.add_slide(blank_layout)
        slide.shapes.add_picture(io.BytesIO(img_bytes), 0, 0, width=prs.slide_width, height=prs.slide_height)
    doc.close()

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def pdf_to_xlsx(data: bytes) -> bytes:
    """
    Extracts tables PyMuPDF can detect per page into worksheet tabs.
    Complex/merged layouts won't be perfect — the API response includes
    a warning flag the frontend surfaces to the user (spec section 12).
    """
    from openpyxl import Workbook

    doc = fitz.open(stream=data, filetype="pdf")
    wb = Workbook()
    wb.remove(wb.active)
    found_any = False

    for i, page in enumerate(doc):
        tables = page.find_tables()
        for t_idx, table in enumerate(tables):
            found_any = True
            ws = wb.create_sheet(title=f"Page{i + 1}_Table{t_idx + 1}"[:31])
            for row in table.extract():
                ws.append(["" if cell is None else cell for cell in row])
    doc.close()

    if not found_any:
        ws = wb.create_sheet(title="Sheet1")
        ws.append(["No tables were detected in this PDF."])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
